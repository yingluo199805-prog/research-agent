# -*- coding: utf-8 -*-
"""通用：读公司三表+解析附注 -> 生成收入分配拆分表 -> 报告勾稽差。
新增公司只需在 COMPANIES 里加一条配置。"""
import sys, os, subprocess, shutil
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from openpyxl.utils import get_column_letter as CL
from gen_revenue_split import build
from parse_notes_generic import parse_company
import common

def yi(v):
    return round(v / 1e8, 5)

def build_company(cfgc):
    path = common.base_path(cfgc['outfile'])
    years = cfgc['years']
    prefix = cfgc['prefix']
    sheet = cfgc['src_sheet']
    # 三表 研发费用行 -> 判定 lump / 未单列
    wb0 = openpyxl.load_workbook(path, data_only=True)
    ws0 = wb0[sheet]
    # 读三表表头, 建 year->列字母映射 (鲁棒, 不依赖 years 是否含最新年)
    import re as _re
    src_year_col = {}
    for rr in range(1, 6):
        for cc in range(2, ws0.max_column + 1):
            v = ws0.cell(rr, cc).value
            if v and _re.match(r'^\s*(\d{4})\s*年报', str(v)):
                src_year_col[int(str(v).strip()[:4])] = CL(cc)
        if src_year_col:
            break
    scol = {y: src_year_col[y] for y in years if y in src_year_col}
    rd_row = cfgc['src_rows']['研发费用']
    def t(r, y):
        v = ws0.cell(r, 2 + (max(years) - y)).value
        return None if v in ('--', None) else v
    RD_T = {y: t(rd_row, y) for y in years}

    notes = parse_company(prefix, years)
    LUMP = "研发费用（未单独披露明细）"

    def collect(sec):
        order = []
        for y in sorted(years, reverse=True):
            rec = notes[y].get(sec)
            if rec:
                for lab in rec['tp']:
                    if lab not in order:
                        order.append(lab)
        return order
    sales_labels = collect('sales'); mgmt_labels = collect('mgmt')
    rd_labels = [LUMP] + collect('rd')
    k = lambda sec, lab: {'sales':'s_','mgmt':'m_','rd':'r_'}[sec] + lab
    sales_items = [(l, k('sales', l)) for l in sales_labels]
    mgmt_items = [(l, k('mgmt', l)) for l in mgmt_labels]
    rd_items = [(l, k('rd', l)) for l in rd_labels]

    warr_cogs = set(cfgc.get('warranty_in_cogs', []))
    warr_vals = cfgc.get('warranty_vals', {})
    note_data = {}
    for y in years:
        d = {}; rec = notes[y]
        s, m, r = rec.get('sales'), rec.get('mgmt'), rec.get('rd')
        if s:
            for lab, v in s['tp'].items(): d[k('sales', lab)] = yi(v)
            if y in warr_cogs: d['售后'] = warr_vals.get(y, 0)
            elif s['after']: d['售后'] = yi(s['after'])
            if s['labor']: d['销售薪酬'] = yi(s['labor'])
            if s['share']: d['销售股份'] = yi(s['share'])
            if s['da']: d['销售折摊'] = yi(s['da'])
        if m:
            for lab, v in m['tp'].items(): d[k('mgmt', lab)] = yi(v)
            if m['labor']: d['管理薪酬'] = yi(m['labor'])
            if m['share']: d['管理股份'] = yi(m['share'])
            if m['da']: d['管理折摊'] = yi(m['da'])
        if r:
            for lab, v in r['tp'].items(): d[k('rd', lab)] = yi(v)
            if r['labor']: d['研发薪酬'] = yi(r['labor'])
            if r['share']: d['研发股份'] = yi(r['share'])
            if r['da']: d['研发折摊'] = yi(r['da'])
        elif RD_T[y] is not None:
            d[k('rd', LUMP)] = round(RD_T[y] / 100, 5)
        note_data[y] = d

    cfg = {
        "sheet_name": "收入分配拆分", "title": cfgc['title'],
        "src_sheet": sheet, "years": years,
        "src_col": scol,
        "unit_div": 100, "warranty_in_cogs": sorted(warr_cogs),
        "mfg_other_pct": 0.009, "src_rows": cfgc['src_rows'],
        "sales_items": sales_items, "mgmt_items": mgmt_items, "rd_items": rd_items,
        "note_data": note_data,
    }
    # ---- 人员数量(可选): cfgc['headcount']={year:{制造人数,销售人数,管理人数,研发人数}} ----
    hc = cfgc.get('headcount')
    if hc:
        for y in years:
            note_data[y].update(hc.get(y, {}))
    # ---- 建附注底稿 + 公式链接 ----
    ws_name = cfgc['title'].replace('收入分配拆分', '') + '附注底稿'
    ws_rows, ws_col = build_worksheet(path, ws_name, note_data, sales_items, mgmt_items,
                                      rd_items, years, bool(hc))
    cfg.update({'ws_sheet': ws_name, 'ws_rows': ws_rows, 'ws_col': ws_col})
    wb = openpyxl.load_workbook(path)
    ws, dr, dlast = build(wb, cfg)
    wb.save(path)
    return path, sales_labels, mgmt_labels, rd_labels

def build_worksheet(path, ws_name, note_data, sales_items, mgmt_items, rd_items, years, has_hc):
    """建『XX附注底稿』sheet（研报三线表样式，可直接导出报告）, 返回 (ws_rows, ws_col)。"""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    NAVY = "0F243E"; GREY = "808080"
    HFONT = Font(name='思源黑体 CN Medium', bold=True, color=NAVY, size=9)   # 表头(参考报告模板)
    SECF = Font(name='思源黑体 CN Medium', bold=True, color=NAVY, size=9)
    BODY = Font(name='思源黑体 CN', size=9)
    BODYB = Font(name='思源黑体 CN', size=9, bold=True)
    TITLE = Font(name='思源黑体 CN Medium', bold=True, color=NAVY, size=11)
    SRC = Font(name='思源黑体 CN', size=8, color=GREY)
    F2 = PatternFill("solid", fgColor="F2F2F2")
    med = Side(style="medium", color=NAVY); thin = Side(style="thin", color="BFBFBF")
    RIGHT = Alignment(horizontal='right'); LEFT = Alignment(horizontal='left'); CEN = Alignment(horizontal='center')
    AMT = "#,##0.00_);[Red](#,##0.00)"
    wb = openpyxl.load_workbook(path)
    if ws_name in wb.sheetnames:
        del wb[ws_name]
    ws = wb.create_sheet(ws_name)
    ws.sheet_view.showGridLines = False
    wcol = {y: 3 + (y - min(years)) for y in years}
    last_col = wcol[years[-1]]
    ws.cell(1, 2, f'表：{ws_name}').font = TITLE
    ws.cell(2, 2, '单位：亿元；人数为人').font = SRC
    HDR = 3
    hc0 = ws.cell(HDR, 2, '科目'); hc0.font = HFONT; hc0.fill = F2; hc0.alignment = LEFT
    for y in years:
        c = ws.cell(HDR, wcol[y], f'{y}年'); c.font = HFONT; c.fill = F2; c.alignment = CEN
    ws_rows = {}; r = [4]
    def sec(t):
        cell = ws.cell(r[0], 2, t); cell.font = SECF; cell.alignment = LEFT
        for c in range(2, last_col + 1):
            ws.cell(r[0], c).fill = F2
        r[0] += 1
    def item(label, key, fmt=AMT):
        lc = ws.cell(r[0], 2, label); lc.font = BODY; lc.alignment = LEFT
        for y in years:
            v = note_data[y].get(key)
            c = ws.cell(r[0], wcol[y], v if v is not None else None)
            c.number_format = fmt; c.font = BODY; c.alignment = RIGHT
        ws_rows[key] = r[0]; r[0] += 1
    sec('一、销售费用明细')
    for lab, key in sales_items: item('—' + lab, key)
    item('售后服务费', '售后'); item('销售人员工资薪金', '销售薪酬')
    item('销售人员股份支付', '销售股份'); item('销售费用中折旧摊销', '销售折摊'); r[0] += 1
    sec('二、管理费用明细')
    for lab, key in mgmt_items: item('—' + lab, key)
    item('管理人员工资薪金', '管理薪酬'); item('管理人员股份支付', '管理股份'); item('管理费用中折旧摊销', '管理折摊'); r[0] += 1
    sec('三、研发费用明细')
    for lab, key in rd_items: item('—' + lab, key)
    item('研发人员工资薪金', '研发薪酬'); item('研发人员股份支付', '研发股份'); item('研发费用中折旧摊销', '研发折摊'); r[0] += 1
    if has_hc:
        sec('四、人员数量（专业构成映射：制造=生产+其他/销售/管理=财务+行政/研发=技术或研发）')
        item('制造工人数量', '制造人数', '#,##0'); item('销售人员数量', '销售人数', '#,##0')
        item('管理人员数量', '管理人数', '#,##0'); item('研发人员数量', '研发人数', '#,##0')
        hcnote = ('注：收入分配表"人均工资"=对应职能费用工资薪金÷此处人数，仅为近似参考。'
                  '费用工资按职能、人数按专业构成，两者口径不完全对应（如管理人数=财务+行政，'
                  '可能小于管理费用工资覆盖面；"其他人员"默认并入制造工人），个别职能人均工资可能偏高/偏低。')
        hn = ws.cell(r[0], 2, hcnote); hn.font = SRC; hn.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r[0], start_column=2, end_row=r[0], end_column=last_col)
        ws.row_dimensions[r[0]].height = 50; r[0] += 1
    # 五、收入勾稽校验（供应商=residual, 加总恒=100%）
    sec('五、收入勾稽校验（供应商为residual, 加总恒=100%）')
    cc = ws.cell(r[0], 2, '加总占比（恒=100%）'); cc.font = BODYB; cc.alignment = LEFT
    for i, y in enumerate(years):
        sc = CL(4 + (y - min(years)))
        c = ws.cell(r[0], wcol[y], f"='收入分配拆分'!{sc}22"); c.number_format = '0.0%'; c.font = BODY; c.alignment = RIGHT
    last_table_row = r[0]
    r[0] += 1
    # ---- 三线表边框 ----
    for c in range(2, last_col + 1):
        ws.cell(HDR, c).border = Border(top=med, bottom=thin)            # 顶线 + 表头下线
        ws.cell(last_table_row, c).border = Border(bottom=med)           # 底线
    # 数据来源 + 备注（表下方, 报告风）
    src = ws.cell(r[0], 2, '数据来源：公司历年年度报告（销售/管理/研发费用、职工薪酬、专业构成附注），当年原始披露口径；广发证券发展研究中心')
    src.font = SRC; ws.merge_cells(start_row=r[0], start_column=2, end_row=r[0], end_column=last_col); r[0] += 1
    note = ('注：供应商=营业总收入−其余各已披露分配项(residual)，即营业成本这一上一级的"其他/给供应商的钱"，'
            '故加总恒=100%、无全局误差项。各费用明细的"其他"行已兜住费用内未单列项(合计对得上三表)。'
            '成本侧口径残差(数据端"营业总成本"约0.3%未拆明细、折旧摊销现金流口径vs成本费用口径差、'
            '职工薪酬应付变动+现金+股份的时点差)自然并入供应商；拆分表附段"口径残差"诊断行已隐藏，需要时取消隐藏可查看其大小。')
    nc = ws.cell(r[0], 2, note); nc.font = SRC; nc.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r[0], start_column=2, end_row=r[0], end_column=last_col)
    ws.row_dimensions[r[0]].height = 76
    ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 32
    from openpyxl.utils import get_column_letter as _CL
    for y in years:
        ws.column_dimensions[_CL(wcol[y])].width = 11
    ws.freeze_panes = 'C4'
    wb.save(path)
    return ws_rows, {y: CL(wcol[y]) for y in years}

def recon(path, years):
    """LibreOffice 重算后读 收入 vs 加总。本机无 LibreOffice 时返回 None。"""
    exe = shutil.which("libreoffice") or shutil.which("soffice")
    if not exe:
        return None
    import tempfile
    d = tempfile.mkdtemp(prefix="rc_")
    tmp = os.path.join(d, "recon.xlsx"); shutil.copy(path, tmp)
    subprocess.run([exe, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(d, "out"), tmp],
                   capture_output=True, timeout=180)
    ws = openpyxl.load_workbook(os.path.join(d, "out", "recon.xlsx"), data_only=True)['收入分配拆分']
    out = []
    for i, y in enumerate(years):
        c = 4 + i; rev = ws.cell(39, c).value; tot = ws.cell(40, c).value
        if rev and tot is not None:
            out.append((y, round((tot-rev)/rev*100, 3)))
    return out


def recon_py(cfgc):
    """不依赖 LibreOffice 的勾稽校验（读三表缓存值 + 解析附注，纯 Python 复算）：
    1) 销售/管理/研发费用附注合计 vs 三表对应行（差 >0.5% 提示解析漏行/口径错）；
    2) 供应商 residual 占收入比 及 成本侧口径残差（=residual 供应商 − 营业成本口径供应商）占收入比。
    加总恒=100% 由 residual 保证，故不再报"加总差"。"""
    path = common.base_path(cfgc['outfile'])
    years = cfgc['years']; R = cfgc['src_rows']; prefix = cfgc['prefix']
    ws = openpyxl.load_workbook(path, data_only=True)[cfgc['src_sheet']]
    import re as _re
    ycol = {}
    for rr in range(1, 6):
        for cc in range(2, ws.max_column + 1):
            v = ws.cell(rr, cc).value
            if v and _re.match(r'^\s*(\d{4})\s*年报', str(v)):
                ycol[int(str(v).strip()[:4])] = cc
        if ycol:
            break
    # 销售/管理/研发费用 行若配置未给，按标签在三表 A 列定位（仅用于校验）
    R = dict(R)
    for lab in ('销售费用', '管理费用', '研发费用'):
        if lab not in R:
            for rr in range(1, ws.max_row + 1):
                v = ws.cell(rr, 1).value
                if v and str(v).replace(' ', '').strip() == lab:
                    R[lab] = rr; break
    def g(key, y, default=0.0):
        r = R.get(key)
        if not r or y not in ycol:
            return default
        v = ws.cell(r, ycol[y]).value
        return float(v) / 100 if isinstance(v, (int, float)) else default   # 百万 -> 亿
    notes = parse_company(prefix, years)
    print("  附注合计 vs 三表(亿元)  [销售 | 管理 | 研发]   供应商%  口径残差%")
    for y in years:
        rec = notes.get(y, {})
        cols = []
        for sec, key in (('sales', '销售费用'), ('mgmt', '管理费用'), ('rd', '研发费用')):
            n = rec.get(sec)
            t = g(key, y, None)
            if n is None or t is None:
                cols.append("   --   ")
            else:
                nt = n['total'] / 1e8
                diff = (nt - t) / t * 100 if t else 0.0
                flag = "" if abs(diff) < 0.5 else "!"
                cols.append(f"{nt:8.2f}/{t:8.2f}{flag}")
        rev = g('营业总收入', y) or g('营业收入', y)
        # 复算供应商 residual：收入 − (经营损失 + 非经营损失 + 业务费用 + 售后 + 员工薪酬 + 银行 + 政府 + 折旧 + 摊销 + 净利)
        op = -(g('其他收益', y) + g('投资收益', y) + g('公允价值变动收益', y) + g('信用减值损失', y)
               + g('资产减值损失', y) + g('资产处置收益', y) + g('汇兑收益', y) + g('净敞口套期损益', y))
        nonop = -(g('营业外收入', y) - g('营业外支出', y))
        bank = g('财务费用', y) + g('利息支出', y) + g('手续费及佣金支出', y)
        gov = g('税金及附加', y) + g('所得税费用', y)
        dep = g('固定资产折旧', y) + g('投资性房地产折旧', y)
        amort = g('无形资产摊销', y) + g('长期待摊费用摊销', y) + g('使用权资产折旧', y)
        ni = g('净利润母公司', y) + g('少数股东损益', y)
        def part(sec, k):
            n = rec.get(sec); return (n[k] / 1e8) if n else 0.0
        tp = sum(sum(v for v in rec[s]['tp'].values()) / 1e8 for s in ('sales', 'mgmt', 'rd') if rec.get(s))
        if not rec.get('rd'):
            tp += g('研发费用', y)            # lump 年：整笔研发进业务费用
        mfg_other = round(rev * 0.009, 2)
        biz = tp + mfg_other
        after = part('sales', 'after')
        share = sum(part(s, 'share') for s in ('sales', 'mgmt', 'rd'))
        prev_ap = g('应付职工薪酬', y - 1, None)
        emp = (g('支付职工现金', y) + ((g('应付职工薪酬', y) - prev_ap) if prev_ap is not None else 0.0)) + share
        supplier = rev - (op + nonop + biz + after + emp + bank + gov + dep + amort + ni)
        indirect_da = sum(part(s, 'da') for s in ('sales', 'mgmt', 'rd'))
        mfg_worker = emp - sum(part(s, 'labor') + part(s, 'share') for s in ('sales', 'mgmt', 'rd'))
        supplier_ref = g('营业成本', y) - mfg_worker - (dep + amort - indirect_da) - mfg_other
        gap = supplier - supplier_ref
        sp = supplier / rev * 100 if rev else 0.0
        gp = gap / rev * 100 if rev else 0.0
        print(f"  {y}: {' | '.join(cols)}   {sp:6.1f}%  {gp:+6.2f}%")
