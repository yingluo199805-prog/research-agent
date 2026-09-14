# -*- coding: utf-8 -*-
"""读取解析得到的附注 JSON, 构建长城 2014-2025 全年份 note_data 并生成收入分配拆分表。"""
import sys, os, json
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from openpyxl.utils import get_column_letter
from gen_revenue_split import build

PATH = "/sessions/lucid-funny-carson/mnt/财报&模型xCowork/长城A股.xlsx"
notes = json.load(open('/tmp/gw_notes.json', encoding='utf-8'))
notes = {int(k): v for k, v in notes.items()}
years = list(range(2014, 2026))

# 三表 row97 研发费用 (百万, 原始披露); None=未单独披露(并入管理)
RD_T = {2025:10432.34,2024:9284.37,2023:8054.25,2022:6445.16,2021:4489.57,
        2020:3067.48,2019:2716.22,2018:1743.38,2017:None,2016:None,2015:None,2014:None}
# 售后(质保)在COGS的年份, 用单独估值
WARR_COGS = {2024: 15.47757, 2025: 17.06}   # 2024=重分类金额; 2025≈按营收比例估
WCOGS_YEARS = set(WARR_COGS)

def yi(v):   # 元 -> 亿元
    return round(v / 1e8, 5)

# ---- 收集各 section 明细科目并集 (recent-first) ----
def collect(sec):
    order = []
    for y in sorted(years, reverse=True):
        rec = notes.get(y, {}).get(sec)
        if not rec:
            continue
        for lab in rec['tp']:
            if lab not in order:
                order.append(lab)
    return order

sales_labels = collect('sales')
mgmt_labels = collect('mgmt')
rd_labels = collect('rd')
# 研发未拆行(并入披露年份用)
LUMP = "研发费用（未单独披露明细）"
rd_labels = [LUMP] + rd_labels

def key(sec, lab):
    return {'sales': 's_', 'mgmt': 'm_', 'rd': 'r_'}[sec] + lab

sales_items = [(l, key('sales', l)) for l in sales_labels]
mgmt_items = [(l, key('mgmt', l)) for l in mgmt_labels]
rd_items = [(l, key('rd', l)) for l in rd_labels]

# ---- 逐年 note_data ----
note_data = {}
for y in years:
    d = {}
    rec = notes.get(y, {})
    s = rec.get('sales'); m = rec.get('mgmt'); r = rec.get('rd')
    if s:
        for lab, v in s['tp'].items():
            d[key('sales', lab)] = yi(v)
        if y in WCOGS_YEARS:
            d['售后'] = WARR_COGS[y]
        elif s['after']:
            d['售后'] = yi(s['after'])
        if s['labor']:
            d['销售薪酬'] = yi(s['labor'])
        if s['share']:
            d['销售股份'] = yi(s['share'])
        if s['da']:
            d['销售折摊'] = yi(s['da'])
    if m:
        for lab, v in m['tp'].items():
            d[key('mgmt', lab)] = yi(v)
        if m['labor']:
            d['管理薪酬'] = yi(m['labor'])
        if m['share']:
            d['管理股份'] = yi(m['share'])
        if m['da']:
            d['管理折摊'] = yi(m['da'])
    if r:   # 详细研发附注
        for lab, v in r['tp'].items():
            d[key('rd', lab)] = yi(v)
        if r['labor']:
            d['研发薪酬'] = yi(r['labor'])
        if r['share']:
            d['研发股份'] = yi(r['share'])
        if r['da']:
            d['研发折摊'] = yi(r['da'])
    elif RD_T[y] is not None:   # 有研发科目但无明细 -> 整笔放未拆行
        d[key('rd', LUMP)] = round(RD_T[y] / 100, 5)
    note_data[y] = d

src_col = {y: get_column_letter(2 + (2025 - y)) for y in years}
cfg = {
    "sheet_name": "收入分配拆分", "title": "长城收入分配拆分",
    "src_sheet": "长城三表", "years": years, "src_col": src_col,
    "unit_div": 100, "warranty_in_cogs": sorted(WCOGS_YEARS),
    "mfg_other_pct": 0.009,
    "src_rows": {"营业总收入":86,"营业收入":87,"营业成本":91,"税金及附加":94,"财务费用":98,
        "利息支出":92,"手续费及佣金支出":93,"资产处置收益":104,"投资收益":102,
        "资产减值损失":105,"信用减值损失":106,"其他收益":107,"公允价值变动收益":101,
        "营业外收入":109,"营业外支出":111,"所得税费用":114,"净利润母公司":117,"少数股东损益":118,
        "支付职工现金":151,"固定资产折旧":195,"投资性房地产折旧":196,"使用权资产折旧":197,
        "无形资产摊销":198,"长期待摊费用摊销":199,"应付职工薪酬":272},
    "sales_items": sales_items, "mgmt_items": mgmt_items, "rd_items": rd_items,
    "note_data": note_data,
}
wb = openpyxl.load_workbook(PATH)
ws, dr, dlast = build(wb, cfg)
wb.save(PATH)
print("OK. sales/mgmt/rd 明细行数:", len(sales_items), len(mgmt_items), len(rd_items))
print("sales labels:", sales_labels)
print("mgmt labels:", mgmt_labels)
print("rd labels:", rd_labels)
