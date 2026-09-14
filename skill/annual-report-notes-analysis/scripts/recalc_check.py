# -*- coding: utf-8 -*-
"""Mac 版重算自检：用 formulas 包重算工作簿（无 Excel/LibreOffice），然后：
① 扫描 #REF!/#DIV/0!/#VALUE!/#NAME?/#N/A；② 打印所有「校验」行的值；③ 打印「折旧率/摊销率/转固率/资本化率/占比」行；
④ 检查预测列（E 年）是否有空白。重算结果另存 <文件>_recalc.xlsx（原文件不动，公式保留）。
用法: python3 recalc_check.py <xlsx> [--sheets 固定资产,无形资产]
"""
import sys, os, re, shutil, tempfile
import openpyxl


def recalc(path):
    import formulas
    d = tempfile.mkdtemp(prefix="rc_")
    xl = formulas.ExcelModel().loads(path).finish()
    xl.calculate()
    xl.write(dirpath=d)
    outs = [os.path.join(d, f) for f in os.listdir(d) if f.lower().endswith(".xlsx")]
    if not outs:
        raise SystemExit("formulas 未输出文件")
    dst = os.path.splitext(path)[0] + "_recalc.xlsx"
    shutil.copy(outs[0], dst)
    return dst


def check(path, sheets=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(os.path.splitext(path)[0].replace("_recalc", "") + ".xlsx", data_only=False)
    ERR = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")
    for name in (sheets or wb.sheetnames):
        ws = wb[name] if name in wb.sheetnames else wb[name.upper()]
        wsf = wbf[name]
        hdr_row = 5
        years = {c: ws.cell(hdr_row, c).value for c in range(4, ws.max_column + 1) if ws.cell(hdr_row, c).value and re.match(r"^\d{4}[AE]$", str(ws.cell(hdr_row, c).value))}
        fc_cols = [c for c, v in years.items() if str(v).endswith("E")]
        print(f"\n=========== {name} ===========")
        errs = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERR:
                    errs.append(c.coordinate)
        print(f"公式错误: {len(errs)} {errs[:12]}")
        blanks = []
        for r in range(6, ws.max_row + 1):
            lab = ws.cell(r, 2).value
            if not lab or (isinstance(lab, str) and (lab.startswith("【") or lab.startswith("一、") or re.match(r"^[一二三四五六七八九十]+[、\-]", lab) or lab.startswith("8.") or lab.startswith("9.") or lab.startswith("6.") or lab.startswith("7."))):
                continue
            for c in fc_cols:
                f = wsf.cell(r, c).value
                v = ws.cell(r, c).value
                if f is not None and (v is None or v == ""):
                    blanks.append(f"{ws.cell(r, c).coordinate}")
        print(f"预测列有公式但算不出值: {len(blanks)} {blanks[:12]}")
        def fmt(v):
            if v is None or v == "": return "   -   "
            if isinstance(v, (int, float)):
                return f"{v:,.0f}" if abs(v) >= 2 else f"{v:.2%}"
            return str(v)[:8]
        print("  " + " | ".join(f"{str(years[c]):>10s}" for c in sorted(years)))
        for r in range(6, ws.max_row + 1):
            lab = ws.cell(r, 2).value
            if not isinstance(lab, str): continue
            if any(k in lab for k in ("校验", "折旧率", "摊销率", "转固率", "资本化率", "新旧度", "占比", "倍数", "费用率", "转出率")):
                vals = [ws.cell(r, c).value for c in sorted(years)]
                flag = ""
                if "校验" in lab and any(isinstance(v, (int, float)) and abs(v) > 0.5 for v in vals):
                    flag = "  <-- 非零"
                print(f"{lab[:38]:38s} " + " | ".join(f"{fmt(v):>10s}" for v in vals) + flag)


if __name__ == "__main__":
    args = sys.argv[1:]
    sheets = None
    if "--sheets" in args:
        i = args.index("--sheets"); sheets = args[i + 1].split(","); del args[i:i + 2]
    src = args[0]
    out = recalc(src)
    print("recalc ->", out)
    check(out, sheets)
