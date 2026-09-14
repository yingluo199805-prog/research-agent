#!/usr/bin/env python
"""Validate formula hygiene and revenue reconciliation for a revenue split sheet."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl


ERROR_TOKENS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A")
NON_AMOUNT_TERMS = (
    "同比",
    "增速",
    "增长率",
    "增长",
    "变动",
    "幅度",
    "占比",
    "比例",
    "利润率",
    "毛利率",
    "净利率",
    "ROE",
    "ROA",
    "周转率",
    "每股",
    "PE",
    "PB",
    "PS",
)
UNIT_MULTIPLIER_TO_YUAN = (
    ("百万元", 1_000_000),
    ("人民币百万元", 1_000_000),
    ("RMB million", 1_000_000),
    ("million", 1_000_000),
    ("亿元", 100_000_000),
    ("人民币亿元", 100_000_000),
    ("万元", 10_000),
    ("人民币万元", 10_000),
    ("千元", 1_000),
    ("人民币千元", 1_000),
    ("元", 1),
    ("人民币元", 1),
)


def referenced_rows(formula: str) -> set[int]:
    rows: set[int] = set()
    for token in re.findall(r"\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?", formula or ""):
        token = token.replace("$", "")
        if ":" in token:
            start, end = token.split(":")
            start_col = re.search(r"[A-Z]+", start).group()
            end_col = re.search(r"[A-Z]+", end).group()
            if start_col != end_col:
                continue
            start_row = int(re.search(r"\d+", start).group())
            end_row = int(re.search(r"\d+", end).group())
            rows.update(range(min(start_row, end_row), max(start_row, end_row) + 1))
        else:
            rows.add(int(re.search(r"\d+", token).group()))
    return rows


def looks_non_amount(label: str, formula: str) -> bool:
    text = f"{label or ''} {formula or ''}".upper()
    if any(term.upper() in text for term in NON_AMOUNT_TERMS):
        return True
    if isinstance(formula, str) and formula.startswith("="):
        compact = formula.replace(" ", "")
        if re.search(r"/[^,()]+-1\)?$", compact):
            return True
    return False


def detect_unit(ws, max_rows: int = 15, max_cols: int = 8):
    text_parts = []
    for row in range(1, min(ws.max_row, max_rows) + 1):
        for col in range(1, min(ws.max_column, max_cols) + 1):
            value = ws.cell(row, col).value
            if value is not None:
                text_parts.append(str(value))
    text = " ".join(text_parts)
    for unit, multiplier in UNIT_MULTIPLIER_TO_YUAN:
        if unit in text:
            return {"unit": unit, "multiplier_to_yuan": multiplier}
    return None


def statement_sheet_for(wb, output_ws):
    candidates = [ws for ws in wb.worksheets if ws.title != output_ws.title and "三表" in ws.title]
    if candidates:
        return candidates[0]
    others = [ws for ws in wb.worksheets if ws.title != output_ws.title]
    return others[0] if others else None


def has_statement_ref(formula: str) -> bool:
    return isinstance(formula, str) and "!" in formula and "三表" in formula


def formula_conversion_factor(formula: str):
    compact = formula.replace(" ", "")
    if re.search(r"[/]100000000(?!\d)", compact):
        return 1 / 100_000_000
    if re.search(r"[/]1000000(?!\d)", compact):
        return 1 / 1_000_000
    if re.search(r"[/]10000(?!\d)", compact):
        return 1 / 10_000
    if re.search(r"[/]1000(?!\d)", compact):
        return 1 / 1_000
    if re.search(r"[/]100(?!\d)", compact):
        return 1 / 100
    if re.search(r"[*]100(?!\d)", compact):
        return 100
    if re.search(r"[*]0\.01(?!\d)", compact):
        return 0.01
    if re.search(r"[*]0\.1(?!\d)", compact):
        return 0.1
    if re.search(r"[*]1000(?!\d)", compact):
        return 1000
    if re.search(r"[*]10000(?!\d)", compact):
        return 10000
    return 1


def close_enough(a: float, b: float) -> bool:
    return abs(a - b) <= max(abs(b) * 1e-9, 1e-12)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--sheet", default="收入拆分")
    parser.add_argument("--output", default="revenue_split_validation.json")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=False)
    ws = wb[args.sheet]
    statement_ws = statement_sheet_for(wb, ws)
    output_unit = detect_unit(ws)
    statement_unit = detect_unit(statement_ws) if statement_ws is not None else None
    expected_factor = None
    if output_unit and statement_unit:
        expected_factor = statement_unit["multiplier_to_yuan"] / output_unit["multiplier_to_yuan"]

    issues = {
        "unit_detection": {
            "statement_sheet": statement_ws.title if statement_ws is not None else None,
            "statement_unit": statement_unit,
            "output_unit": output_unit,
            "expected_statement_to_output_factor": expected_factor,
        },
        "unit_conversion_mismatch": [],
        "unquoted_source_sheet": [],
        "raw_source_ref_not_wrapped_in_N": [],
        "formula_error_tokens": [],
        "non_amount_rows_in_revenue_bridge": [],
        "view_or_hyperlink_issues": [],
        "revenue_checks": [],
    }

    if ws.freeze_panes:
        freeze = str(ws.freeze_panes)
        match = re.search(r"(\d+)", freeze)
        freeze_row = int(match.group(1)) if match else 1
        if freeze_row > 10:
            issues["view_or_hyperlink_issues"].append(
                {
                    "type": "deep_freeze_pane",
                    "freeze_panes": freeze,
                    "message": "Deep freeze panes can make generated sheets appear unable to scroll; reset or justify it.",
                }
            )

    for coord, expected_targets in {"A1": ("C3", "B3"), "A2": ("C37", "B37")}.items():
        link = ws[coord].hyperlink
        if not link:
            issues["view_or_hyperlink_issues"].append(
                {"type": "missing_navigation_hyperlink", "cell": coord}
            )
            continue
        location = link.location or link.target or ""
        if "Sheet1!" in location or "Sheet1" in location:
            issues["view_or_hyperlink_issues"].append(
                {"type": "stale_sheet1_hyperlink", "cell": coord, "location": location}
            )
        if ws.title not in location:
            issues["view_or_hyperlink_issues"].append(
                {"type": "wrong_sheet_hyperlink", "cell": coord, "location": location, "expected_sheet": ws.title}
            )
        if not any(target in location for target in expected_targets):
            issues["view_or_hyperlink_issues"].append(
                {"type": "unexpected_navigation_target", "cell": coord, "location": location, "expected_any": expected_targets}
            )

    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if not (isinstance(value, str) and value.startswith("=")):
                continue
            coord = cell.coordinate
            if has_statement_ref(value) and expected_factor is not None:
                actual_factor = formula_conversion_factor(value)
                if not close_enough(actual_factor, expected_factor):
                    issues["unit_conversion_mismatch"].append(
                        {
                            "cell": coord,
                            "formula": value,
                            "actual_detected_factor": actual_factor,
                            "expected_factor": expected_factor,
                        }
                    )
            if re.search(r"(?<!')\b[^']*三表!", value):
                issues["unquoted_source_sheet"].append(coord)
            if "'长城三表'!" in value and "N('长城三表'!" not in value:
                issues["raw_source_ref_not_wrapped_in_N"].append(coord)
            if any(token in value for token in ERROR_TOKENS):
                issues["formula_error_tokens"].append(coord)

    # Inspect formulas whose labels indicate the revenue split total. Any rows
    # referenced by these formulas must be monetary amount rows, not growth/KPI rows.
    for row in ws.iter_rows():
        row_num = row[0].row
        row_label = " ".join(str(ws.cell(row_num, col).value or "") for col in range(1, min(ws.max_column, 4) + 1))
        if "下面各项加总" not in row_label and "收入拆分加总" not in row_label:
            continue
        nearby_context = " ".join(
            str(ws.cell(r, col).value or "")
            for r in range(max(1, row_num - 5), row_num + 1)
            for col in range(1, min(ws.max_column, 4) + 1)
        )
        if "项目占比" in nearby_context:
            continue
        for cell in row:
            value = cell.value
            if not (isinstance(value, str) and value.startswith("=") and "SUM(" in value.upper()):
                continue
            for ref_row in sorted(referenced_rows(value)):
                label = " ".join(str(ws.cell(ref_row, col).value or "") for col in range(1, min(ws.max_column, 4) + 1))
                formula = ws.cell(ref_row, cell.column).value
                if looks_non_amount(label, formula):
                    issues["non_amount_rows_in_revenue_bridge"].append(
                        {
                            "total_cell": cell.coordinate,
                            "referenced_row": ref_row,
                            "label": label,
                            "same_column_formula": formula,
                        }
                    )

    # If the standard check block exists, capture its formulas for review.
    for row in range(177, min(ws.max_row, 190) + 1):
        label = ws.cell(row, 2).value or ws.cell(row, 3).value
        if label:
            issues["revenue_checks"].append({
                "row": row,
                "label": label,
                "formulas": [ws.cell(row, col).value for col in range(4, min(ws.max_column, 15) + 1)],
            })

    Path(args.output).write_text(json.dumps(issues, ensure_ascii=False, indent=2), encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
