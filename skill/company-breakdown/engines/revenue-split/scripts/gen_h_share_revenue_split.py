# -*- coding: utf-8 -*-
"""H股收入分配拆分生成器。

适用于 H 股年报按性质披露、且附注项已抽成 note_data 的场景。
保留比亚迪/A股模型的三段结构：总项统计 / 项目占比 / 细项拆分。
"""
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter as CL


NAVY = "0F243E"
HDR_BLUE = "2E3160"
SEC_FILL = "E7E9F0"
LIGHT_GREY = "F2F2F2"
GREY = "808080"
AMT = "#,##0.00_);[Red](#,##0.00)"
PCT = "0.0%"


def detect_year_columns(ws, years):
    """Return {year: column_letter} by scanning the top rows for 'YYYY 年报'."""
    out = {}
    for r in range(1, min(ws.max_row, 8) + 1):
        for c in range(2, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not v:
                continue
            s = str(v).strip()
            if len(s) >= 4 and s[:4].isdigit() and "年报" in s:
                y = int(s[:4])
                if y in years:
                    out[y] = CL(c)
        if len(out) == len(years):
            break
    missing = [y for y in years if y not in out]
    if missing:
        raise ValueError(f"Cannot locate year columns in source sheet: {missing}")
    return out


def build_note_sheet(wb, cfg):
    """Create the H-share note working paper and return (row_map, col_map)."""
    name = cfg["note_sheet"]
    years = cfg["years"]
    rows = cfg["note_rows"]
    data = cfg["note_data"]
    unit_label = cfg.get("note_unit_label", "单位：人民币百万元")

    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ycol = {y: 3 + i for i, y in enumerate(years)}
    last_col = ycol[years[-1]]

    title_font = Font(name="思源黑体 CN Medium", bold=True, color=NAVY, size=11)
    src_font = Font(name="思源黑体 CN", size=8, color=GREY)
    head_font = Font(name="思源黑体 CN Medium", bold=True, color=NAVY, size=9)
    body_font = Font(name="思源黑体 CN", size=9)
    sec_font = Font(name="思源黑体 CN Medium", bold=True, color=NAVY, size=9)
    fill = PatternFill("solid", fgColor=LIGHT_GREY)
    med = Side(style="medium", color=NAVY)
    thin = Side(style="thin", color="BFBFBF")

    ws.cell(1, 2, cfg.get("note_title", f"表：{name}")).font = title_font
    ws.cell(2, 2, unit_label).font = src_font
    ws.cell(3, 2, "科目").font = head_font
    ws.cell(3, 2).fill = fill
    for y in years:
        cell = ws.cell(3, ycol[y], f"{y}年")
        cell.font = head_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    row_map = {}
    r = 4
    for row in rows:
        typ = row.get("type", "item")
        if typ == "blank":
            r += 1
            continue
        if typ == "section":
            ws.cell(r, 2, row["label"]).font = sec_font
            for c in range(2, last_col + 1):
                ws.cell(r, c).fill = fill
            r += 1
            continue
        if typ == "note":
            cell = ws.cell(r, 2, row["label"])
            cell.font = src_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last_col)
            ws.row_dimensions[r].height = row.get("height", 48)
            r += 1
            continue

        key = row["key"]
        ws.cell(r, 2, row["label"]).font = body_font
        ws.cell(r, 2).alignment = Alignment(horizontal="left")
        for y in years:
            v = data.get(y, {}).get(key)
            cell = ws.cell(r, ycol[y], v)
            cell.number_format = row.get("number_format", "#,##0.000")
            cell.font = body_font
            cell.alignment = Alignment(horizontal="right")
        row_map[key] = r
        r += 1

    last_table_row = max(3, r - 1)
    for c in range(2, last_col + 1):
        ws.cell(3, c).border = Border(top=med, bottom=thin)
        ws.cell(last_table_row, c).border = Border(bottom=med)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 64
    for y in years:
        ws.column_dimensions[CL(ycol[y])].width = 12
    ws.freeze_panes = "C4"
    return row_map, {y: CL(ycol[y]) for y in years}


def build_h_share(wb, cfg):
    """Build H-share revenue split sheet. cfg values are in source/report units."""
    sheet_name = cfg.get("sheet_name", "收入分配拆分")
    years = cfg["years"]
    src_sheet = cfg["src_sheet"]
    src = wb[src_sheet]
    src_col = cfg.get("src_col") or detect_year_columns(src, years)
    src_rows = cfg["src_rows"]
    src_div = cfg.get("src_unit_div", 100)       # 三表百万元 -> 亿元
    note_div = cfg.get("note_unit_div", 100)     # 附注百万元 -> 亿元

    note_rows = cfg.get("note_rows")
    if note_rows:
        note_row, note_col = build_note_sheet(wb, cfg)
    else:
        note_row, note_col = cfg["note_row"], cfg["note_col"]

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    def sref(y, key):
        row = src_rows.get(key)
        if not row:
            return "0"
        return f"IFERROR('{src_sheet}'!{src_col[y]}{row}/{src_div},0)"

    def nref(y, key):
        row = note_row.get(key)
        if not row:
            return "0"
        return f"IFERROR('{cfg['note_sheet']}'!{note_col[y]}{row}/{note_div},0)"

    yc = {y: 4 + i for i, y in enumerate(years)}
    last_col = yc[years[-1]]
    detail = []
    dr = {}

    def add(label, remark=None, group=None, formula=None, key=None):
        row = 39 + len(detail)
        detail.append((row, group, remark, label, formula))
        if key or label:
            dr[key or label] = row
        return row

    add("收入", "营业额/收入", formula=lambda y, L: "=" + sref(y, "收入"))
    add("下面各项加总", formula=lambda y, L: "=" + "+".join(f"{L}{dr[k]}" for k in [
        "经营损失", "非经营损失", "业务费用（给广泛第三方的钱）", "售后服务费（原隶属于销售费用，未单列）",
        "员工薪酬", "供应商（含委外加工及未拆分费用，residual）", "银行", "政府", "折旧", "摊销",
        "净利润-少数股东", "净利润-母公司股东"]))

    op_items = cfg.get("operating_items", [
        ("—其他收入/收益净额", "其他收入", "src"),
        ("—融资收入（利息收入）", "融资收入", "note"),
        ("—经营溢利特殊科目", "经营特殊科目", "src"),
        ("—应占联营公司业绩", "联营", "src"),
        ("—应占合营/共同控制实体业绩", "合营", "src"),
        ("—税前利润特殊科目", "税前特殊", "src"),
    ])
    op_start = add("经营损失", "其他收入+融资收入+联营/合营/特殊科目损益", formula=None)
    first_child = 39 + len(detail)
    for label, key, source in op_items:
        add(label, formula=(lambda k=key, src_type=source: (lambda y, L: "=" + (nref(y, k) if src_type == "note" else sref(y, k))))())
    last_child = 39 + len(detail) - 1
    detail[op_start - 39] = (op_start, None, "其他收入+融资收入+联营/合营/特殊科目损益", "经营损失",
                            lambda y, L: f"=-SUM({L}{first_child}:{L}{last_child})")

    add("非经营损失", "H股通常无独立营业外收支；如披露则在配置中替换", formula=lambda y, L: "=0")
    biz_mode = cfg.get("business_mode", "sum_items")
    biz_remark = cfg.get("business_remark", "按已披露第三方业务费用填列；不能从功能费用中剔除人工/折摊时不强拆")
    biz_start = add("业务费用（给广泛第三方的钱）", biz_remark, "固定成本")
    biz_first = 39 + len(detail)
    biz_items = cfg.get("business_items", [
        ("—研发费用/研发成本（费用化部分）", "研发成本"),
        ("—销售及分销/销售费用（未拆品类明细时整额）", "销售费用"),
        ("—行政/管理费用（未拆品类明细时整额）", "管理费用"),
    ])
    for label, key in biz_items:
        add(label, formula=lambda y, L, k=key: "=" + nref(y, k))
    biz_last = 39 + len(detail) - 1
    if biz_mode == "functional_residual":
        biz_formula = lambda y, L: (
            f"=MAX(SUM({L}{biz_first}:{L}{biz_last})"
            f"-{L}{dr['员工薪酬']}-{L}{dr['折旧']}-{L}{dr['摊销']},0)"
        )
    elif biz_mode == "disclosed_third_party":
        biz_formula = lambda y, L: "=" + nref(y, "业务费用第三方")
    else:
        biz_formula = lambda y, L: f"=SUM({L}{biz_first}:{L}{biz_last})"
    detail[biz_start - 39] = (biz_start, "固定成本", biz_remark, "业务费用（给广泛第三方的钱）", biz_formula)

    for label, key in cfg.get("business_reference_items", []):
        add(label, "参考：功能费用整额，不进入业务费用加总", formula=lambda y, L, k=key: "=" + nref(y, k))

    add("售后服务费（原隶属于销售费用，未单列）", cfg.get("warranty_note", "未单独披露则为0"), "固定成本",
        formula=lambda y, L: "=" + nref(y, "售后"))
    add("员工薪酬", "雇员福利开支总额，通常含股份支付", None, formula=lambda y, L: "=" + nref(y, "员工合计"))
    for label, key in cfg.get("employee_items", [
        ("—薪金、工资及其他福利", "薪金福利"),
        ("—退休福利计划供款", "退休供款"),
        ("—股份支付", "股份支付"),
    ]):
        add(label, formula=lambda y, L, k=key: "=" + nref(y, k))
    add("供应商（含委外加工及未拆分费用，residual）",
        "收入-其余已识别桶；吸收COGS及未拆SG&A中的材料、外协、广告、物流等", "可变成本",
        formula=lambda y, L: f"={L}{dr['收入']}-(" + "+".join(f"{L}{dr[k]}" for k in [
            "经营损失", "非经营损失", "业务费用（给广泛第三方的钱）", "售后服务费（原隶属于销售费用，未单列）",
            "员工薪酬", "银行", "政府", "折旧", "摊销", "净利润-少数股东", "净利润-母公司股东"]) + ")")
    add("银行", "融资成本(gross，未扣融资收入)", "固定成本", formula=lambda y, L: "=" + nref(y, "融资成本"))
    add("政府", "所得税费用；H股通常无A股式税金及附加", "固定成本", formula=lambda y, L: "=" + sref(y, "所得税"))
    add("折旧", "物业、厂房及设备折旧", "固定成本", formula=lambda y, L: "=" + nref(y, "折旧"))
    amort = add("摊销", None, "固定成本")
    amort_first = 39 + len(detail)
    for label, key in cfg.get("amort_items", [
        ("—使用权资产折旧/摊销", "使用权资产折旧"),
        ("—土地使用权/预付土地租赁摊销", "土地摊销"),
        ("—无形资产摊销", "无形资产摊销"),
    ]):
        add(label, formula=lambda y, L, k=key: "=" + nref(y, k))
    amort_last = 39 + len(detail) - 1
    detail[amort - 39] = (amort, "固定成本", None, "摊销", lambda y, L: f"=SUM({L}{amort_first}:{L}{amort_last})")
    add("净利润-少数股东", "非控股权益/少数股东损益", None, formula=lambda y, L: "=" + sref(y, "少数股东损益"))
    add("净利润-母公司股东", "归母净利润", None, formula=lambda y, L: "=" + sref(y, "归母净利润"))
    detail_last = detail[-1][0]

    # Headers and top sections.
    ws.cell(1, 1, "跳转：总项统计")
    ws.cell(1, 1).hyperlink = Hyperlink(ref="A1", location=f"'{sheet_name}'!C3")
    ws.cell(1, 2, "单位：亿元")
    ws.cell(2, 1, "跳转：细项拆分")
    ws.cell(2, 1).hyperlink = Hyperlink(ref="A2", location=f"'{sheet_name}'!C37")
    ws.cell(2, 3, cfg["title"])
    for y in years:
        ws.cell(2, yc[y], f"{y}年")
    for row, title in [(3, "总项统计"), (20, "项目占比"), (37, "细项拆分")]:
        ws.cell(row, 3, title)
        for c in range(3, last_col + 1):
            ws.cell(row, c).fill = PatternFill("solid", fgColor=SEC_FILL)
            ws.cell(row, c).font = Font(name="思源黑体 CN Medium", bold=True, color=NAVY, size=9)
    for row in [4, 38]:
        ws.cell(row, 3, "项目数额")
        for y in years:
            ws.cell(row, yc[y], f"{y}年")
    ws.cell(20, 3, "项目占比")
    for y in years:
        ws.cell(20, yc[y], f"{y}年")

    bridge = ["收入", "下面各项加总", "经营损失", "非经营损失", "业务费用（给广泛第三方的钱）",
              "售后服务费（原隶属于销售费用，未单列）", "员工薪酬",
              "供应商（含委外加工及未拆分费用，residual）", "银行", "政府", "折旧", "摊销",
              "净利润-少数股东", "净利润-母公司股东"]
    for i, label in enumerate(bridge, start=5):
        ws.cell(i, 3, label)
        for y in years:
            idx = years.index(y) + 2
            ws.cell(i, yc[y], f'=VLOOKUP($C{i},$C$39:${CL(last_col)}${detail_last},{idx},FALSE())')
    for i, label in enumerate(bridge, start=21):
        ws.cell(i, 3, label)
        for y in years:
            L = CL(yc[y])
            ws.cell(i, yc[y], 1 if label == "收入" else (f"=SUM({L}23:{L}34)" if label == "下面各项加总" else f"={L}{i-16}/{L}$5"))
            ws.cell(i, yc[y]).number_format = PCT

    for row, group, remark, label, formula in detail:
        ws.cell(row, 1, None)
        ws.cell(row, 2, remark)
        ws.cell(row, 3, label)
        for y in years:
            L = CL(yc[y])
            ws.cell(row, yc[y], formula(y, L) if formula else None)
            ws.cell(row, yc[y]).number_format = AMT

    pct_start = detail_last + 3
    ws.cell(pct_start, 3, "细项占比")
    for y in years:
        ws.cell(pct_start, yc[y], f"{y}年")
    for row, _, _, label, _ in detail:
        out_row = pct_start + (row - 38)
        ws.cell(out_row, 3, label)
        for y in years:
            L = CL(yc[y])
            ws.cell(out_row, yc[y], f"={L}{row}/{L}$39")
            ws.cell(out_row, yc[y]).number_format = PCT
    src_note_row = pct_start + (detail_last - 38) + 2
    ws.cell(src_note_row, 3, "数据来源：公司历年年度报告及三大财务报表、广发证券发展研究中心")

    # Styling.
    body_font = Font(name="思源黑体 CN", size=9)
    bold_font = Font(name="思源黑体 CN", size=9, bold=True)
    head_font = Font(name="思源黑体 CN Medium", size=9, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor=HDR_BLUE)
    sec_fill = PatternFill("solid", fgColor=SEC_FILL)
    thin = Side(style="thin", color="BFBFBF")
    med = Side(style="medium", color=NAVY)
    for r in range(1, src_note_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.font = body_font
            cell.alignment = Alignment(horizontal="right" if c >= 4 else "left", vertical="center")
    for cell in [ws["A1"], ws["A2"]]:
        cell.font = Font(name="思源黑体 CN", size=9, color="0563C1", underline="single")
    for r in [4, 20, 38, pct_start]:
        for c in range(3, last_col + 1):
            ws.cell(r, c).fill = hdr_fill
            ws.cell(r, c).font = head_font
            ws.cell(r, c).alignment = Alignment(horizontal="center")
    for r in [3, 37]:
        for c in range(3, last_col + 1):
            ws.cell(r, c).fill = sec_fill
            ws.cell(r, c).font = bold_font
    bridge_labels = set(bridge)
    detail_summary_rows = [row for row, _, _, label, _ in detail if label in bridge_labels]
    for r in list(range(5, 19)) + list(range(21, 35)) + detail_summary_rows:
        ws.cell(r, 3).font = bold_font
    for r in detail_summary_rows:
        for c in range(3, last_col + 1):
            ws.cell(r, c).fill = sec_fill
            ws.cell(r, c).font = bold_font
    for row, _, _, label, _ in detail:
        pct_row = pct_start + (row - 38)
        if label in bridge_labels:
            for c in range(3, last_col + 1):
                ws.cell(pct_row, c).fill = sec_fill
                ws.cell(pct_row, c).font = bold_font
    for c in range(3, last_col + 1):
        for r in [4, 20, 38, pct_start]:
            ws.cell(r, c).border = Border(top=med, bottom=thin)
        ws.cell(src_note_row, c).border = Border(bottom=med)
    for col, width in [("A", 12), ("B", 42), ("C", 34)]:
        ws.column_dimensions[col].width = width
    for y in years:
        ws.column_dimensions[CL(yc[y])].width = 12
    ws.freeze_panes = None
    return ws
