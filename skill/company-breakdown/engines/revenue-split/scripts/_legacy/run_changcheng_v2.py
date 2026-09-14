# -*- coding: utf-8 -*-
"""长城 v2：建『长城附注底稿』sheet -> 收入分配拆分表公式链接底稿 + 人员数量/人均工资 + 统一格式。"""
import sys, os, json
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from gen_revenue_split import build, AMT, BOLD, SEC_FILL, HDR_FILL

PATH = "/sessions/lucid-funny-carson/mnt/财报&模型xCowork/长城A股.xlsx"
notes = {int(k): v for k, v in json.load(open('/tmp/gw_notes.json', encoding='utf-8')).items()}
years = list(range(2014, 2026))

# 人员数量(生产,销售,技术,财务,行政,其他) 来自年报专业构成
HC = {2014:(33128,1494,10767,846,5841,0),2015:(33124,1070,16512,773,5435,0),
 2016:(38373,1474,10236,741,4450,16343),2017:(34238,1548,17917,812,5079,8911),
 2018:(30575,1722,16972,852,4324,9010),2019:(26675,2019,17603,899,3240,9320),
 2020:(26008,2157,19347,1014,3396,11252),2021:(31363,3283,21137,1688,5994,14469),
 2022:(35329,4254,23850,2106,6022,15806),2023:(36261,3856,24157,1956,5050,11159),
 2024:(37375,6189,23308,1907,4796,11340),2025:(47441,10860,27097,2099,5052,5051)}
def hc(y):
    p,s,t,f,a,o = HC[y]
    return {'制造人数':p+o, '销售人数':s, '管理人数':f+a, '研发人数':t}

RD_T = {2025:10432.34,2024:9284.37,2023:8054.25,2022:6445.16,2021:4489.57,2020:3067.48,
        2019:2716.22,2018:1743.38,2017:None,2016:None,2015:None,2014:None}
WARR = {2024:15.47757, 2025:17.06}
LUMP = "研发费用（未单独披露明细）"
yi = lambda v: round(v/1e8, 5)

def collect(sec):
    order=[]
    for y in sorted(years, reverse=True):
        rec=notes[y].get(sec)
        if rec:
            for lab in rec['tp']:
                if lab not in order: order.append(lab)
    return order
sales_l, mgmt_l = collect('sales'), collect('mgmt')
rd_l = [LUMP] + collect('rd')
k = lambda sec,lab: {'sales':'s_','mgmt':'m_','rd':'r_'}[sec]+lab
sales_items=[(l,k('sales',l)) for l in sales_l]
mgmt_items=[(l,k('mgmt',l)) for l in mgmt_l]
rd_items=[(l,k('rd',l)) for l in rd_l]

# ---- per-year 数据(亿元) ----
note_data={}
for y in years:
    d={}; rec=notes[y]; s,m,r=rec.get('sales'),rec.get('mgmt'),rec.get('rd')
    if s:
        for lab,v in s['tp'].items(): d[k('sales',lab)]=yi(v)
        if y in WARR: d['售后']=WARR[y]
        elif s['after']: d['售后']=yi(s['after'])
        if s['labor']: d['销售薪酬']=yi(s['labor'])
        if s['share']: d['销售股份']=yi(s['share'])
        if s['da']: d['销售折摊']=yi(s['da'])
    if m:
        for lab,v in m['tp'].items(): d[k('mgmt',lab)]=yi(v)
        if m['labor']: d['管理薪酬']=yi(m['labor'])
        if m['share']: d['管理股份']=yi(m['share'])
        if m['da']: d['管理折摊']=yi(m['da'])
    if r:
        for lab,v in r['tp'].items(): d[k('rd',lab)]=yi(v)
        if r['labor']: d['研发薪酬']=yi(r['labor'])
        if r['share']: d['研发股份']=yi(r['share'])
        if r['da']: d['研发折摊']=yi(r['da'])
    elif RD_T[y] is not None:
        d[k('rd',LUMP)]=round(RD_T[y]/100,5)
    d.update(hc(y))
    note_data[y]=d

# ================= 建『长城附注底稿』 =================
wb = openpyxl.load_workbook(PATH)
WS='长城附注底稿'
if WS in wb.sheetnames: del wb[WS]
ws=wb.create_sheet(WS)
wcol={y:3+(y-2014) for y in years}   # C=2014..N=2025
ws.cell(1,1,'长城附注底稿（单位：亿元；人数为人）').font=BOLD
ws.cell(2,1,'数据来源：长城各年度年报附注（销售/管理/研发费用、职工薪酬、专业构成）；当年原始披露口径').font=Font(italic=True,size=9)
ws.cell(3,2,'科目').font=BOLD
for y in years: ws.cell(3,wcol[y],f'{y}年').font=BOLD
ws_rows={}
r=4
def sec_hdr(title):
    global r
    ws.cell(r,1,title).font=BOLD; ws.cell(r,1).fill=SEC_FILL; r+=1
def item(label,key,fmt=AMT):
    global r
    ws.cell(r,2,label)
    for y in years:
        v=note_data[y].get(key)
        c=ws.cell(r,wcol[y], v if v is not None else None); c.number_format=fmt
    ws_rows[key]=r; r+=1

sec_hdr('一、销售费用明细')
for lab,key in sales_items: item('—'+lab,key)
item('售后服务费','售后')
item('销售人员工资薪金','销售薪酬'); item('销售人员股份支付','销售股份'); item('销售费用中折旧摊销','销售折摊')
r+=1
sec_hdr('二、管理费用明细')
for lab,key in mgmt_items: item('—'+lab,key)
item('管理人员工资薪金','管理薪酬'); item('管理人员股份支付','管理股份'); item('管理费用中折旧摊销','管理折摊')
r+=1
sec_hdr('三、研发费用明细')
for lab,key in rd_items: item('—'+lab,key)
item('研发人员工资薪金','研发薪酬'); item('研发人员股份支付','研发股份'); item('研发费用中折旧摊销','研发折摊')
r+=1
sec_hdr('四、人员数量（专业构成→映射）')
item('制造工人数量（生产+其他）','制造人数',fmt='#,##0')
item('销售人员数量','销售人数',fmt='#,##0')
item('管理人员数量（财务+行政）','管理人数',fmt='#,##0')
item('研发人员数量（技术）','研发人数',fmt='#,##0')
ws.column_dimensions['A'].width=22; ws.column_dimensions['B'].width=30
for y in years: ws.column_dimensions[CL(wcol[y])].width=11
ws.freeze_panes='C4'
wb.save(PATH)

# ================= 建收入分配拆分(链接底稿) =================
cfg={'sheet_name':'收入分配拆分','title':'长城收入分配拆分','src_sheet':'长城三表',
 'years':years,'src_col':{y:CL(2+(2025-y)) for y in years},'unit_div':100,
 'warranty_in_cogs':[2024,2025],'mfg_other_pct':0.009,
 'ws_sheet':WS,'ws_rows':ws_rows,'ws_col':{y:CL(wcol[y]) for y in years},
 'src_rows':{'营业总收入':86,'营业收入':87,'营业成本':91,'税金及附加':94,'财务费用':98,'投资收益':102,
   '资产减值损失':105,'信用减值损失':106,'其他收益':107,'公允价值变动收益':101,'营业外收入':109,
   '营业外支出':111,'所得税费用':114,'净利润母公司':117,'少数股东损益':118,'支付职工现金':151,
   '固定资产折旧':195,'投资性房地产折旧':196,'使用权资产折旧':197,'无形资产摊销':198,
   '长期待摊费用摊销':199,'应付职工薪酬':272},
 'sales_items':sales_items,'mgmt_items':mgmt_items,'rd_items':rd_items,'note_data':note_data}
wb=openpyxl.load_workbook(PATH)
ws2,dr,dlast=build(wb,cfg)
wb.save(PATH)
print('OK. 底稿 rows:',len(ws_rows),' split detail_last:',dlast)
