# -*- coding: utf-8 -*-
"""
通用「收入分配拆分」表生成器 —— build-revenue-split skill 核心脚本。
复刻比亚迪模型『XX收入分配拆分』的三段结构(总项统计 / 项目占比 / 细项拆分)与桥接逻辑。
14 个固定分配桥接科目；销售/管理/研发『业务费用』明细行按各公司实际披露科目可变。
单位换算、售后(质保)在COGS的年度差异、A股研发披露演进 等按 cfg 参数化。
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as CL

THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
TOT_FILL = PatternFill("solid", fgColor="F2F2F2")
SEC_FILL = PatternFill("solid", fgColor="BDD7EE")
AMT = "#,##0.00_);[Red](#,##0.00)"
PCT = "0.0%"
BOLD = Font(bold=True)


def build(wb, cfg):
    name = cfg["sheet_name"]
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    src = cfg["src_sheet"]
    years = cfg["years"]
    sc = cfg["src_col"]           # year -> 三表列字母
    u = cfg["unit_div"]           # 百万->亿 = 100
    R = cfg["src_rows"]
    nd = cfg["note_data"]         # year -> {key: 亿元值}
    sales = cfg["sales_items"]    # [label,...]
    mgmt = cfg["mgmt_items"]
    rd = cfg["rd_items"]
    wcogs = set(cfg["warranty_in_cogs"])
    title = cfg["title"]

    c0 = 4
    yc = {y: c0 + i for i, y in enumerate(years)}   # year -> 输出列idx
    cols = [yc[y] for y in years]

    def ref(y, row):
        return f"IFERROR({src}!{sc[y]}{row}/{u},0)"

    def prevcol(y):
        return sc.get(y - 1)

    # ============ 细项拆分区 (从 base 行开始) ============
    base = 39
    r = base
    dr = {}          # 分配科目名 -> 行号
    layout = []      # (row, A, B, Clabel, kind, payload)

    def add(A, B, C, kind, payload=None, key=None):
        nonlocal r
        layout.append((r, A, B, C, kind, payload))
        if key:
            dr[key] = r
        r += 1
        return r - 1

    add(None, "营业总收入", "收入", "rev", R.get("营业总收入", R["营业收入"]), key="收入")
    add(None, None, "下面各项加总", "grand", None, key="加总")
    # 经营损失 (含资产处置收益, 若披露)
    r_op = add(None, "其他经营收益", "经营损失", "negsum", None, key="经营损失")
    op_kids = [("—其他收益", R["其他收益"]), ("—投资收益", R["投资收益"]),
               ("—公允价值变动收益", R["公允价值变动收益"]),
               ("—信用减值损失", R["信用减值损失"]), ("—资产减值损失", R["资产减值损失"])]
    if R.get("资产处置收益"):
        op_kids.append(("—资产处置收益", R["资产处置收益"]))
    if R.get("汇兑收益"):
        op_kids.append(("—汇兑收益", R["汇兑收益"]))
    if R.get("净敞口套期损益"):
        op_kids.append(("—净敞口套期损益", R["净敞口套期损益"]))
    op_first = r
    for lbl, rr in op_kids:
        add(None, None, lbl, "ref", rr)
    op_last = r - 1
    # 非经营损失
    add(None, "营业外收入-营业外支出", "非经营损失", "nonop", None, key="非经营损失")
    # 业务费用
    add("固定成本", None, "业务费用（给广泛第三方的钱）", "biz", None, key="业务费用")
    r_sales_sum = add("固定成本", None, "销售相关费用", "sum_sales", None, key="销售相关费用")
    sales_first = r
    for lbl, dk in sales:
        add(None, None, "—" + lbl, "data", dk)
    sales_last = r - 1
    # 售后服务费 (独立桥接科目)
    add("固定成本", None, "售后服务费（给客户预备的钱，原隶属于销售费用）", "data", "售后", key="售后服务费")
    r_mgmt_sum = add(None, None, "管理相关费用", "sum_mgmt", None, key="管理相关费用")
    mgmt_first = r
    for lbl, dk in mgmt:
        add(None, None, "—" + lbl, "data", dk)
    mgmt_last = r - 1
    r_rd_sum = add(None, None, "研发相关费用", "sum_rd", None, key="研发相关费用")
    rd_first = r
    for lbl, dk in rd:
        add(None, None, "—" + lbl, "data", dk)
    rd_last = r - 1
    add(None, None, "制造费用中除人工、折旧、摊销以外的费用（例如水电等）", "data", "制造其他", key="制造其他")
    # 员工薪酬
    add(None, "Δ应付职工薪酬+支付给职工现金+销管研股份支付", "员工薪酬", "emp", None, key="员工薪酬")
    add("固定成本", "营业成本直接人工+制造费用人工", "制造工人薪酬（营业成本中的直接人工+制造费用中的人工）", "mfgworker", None, key="制造工人")
    add(None, None, "制造工人数量", "data", "制造人数")
    add(None, None, "—制造人员工资", "wage", "制造")
    r_semp = add("固定成本", None, "销售人员职工薪酬", "sum2", None, key="销售人员")
    add(None, None, "—销售人员职工薪酬", "data", "销售薪酬")
    add(None, None, "—销售人员股份支付", "data", "销售股份")
    add(None, None, "—销售人员数量", "data", "销售人数")
    add(None, None, "—销售人员工资", "wage", "销售")
    r_memp = add("固定成本", None, "管理人员职工薪酬", "sum2", None, key="管理人员")
    add(None, None, "—管理人员职工薪酬", "data", "管理薪酬")
    add(None, None, "—管理人员股份支付", "data", "管理股份")
    add(None, None, "—管理人员数量", "data", "管理人数")
    add(None, None, "—管理人员工资", "wage", "管理")
    r_remp = add("固定成本", None, "研发人员职工薪酬", "sum2", None, key="研发人员")
    add(None, None, "—研发人员职工薪酬", "data", "研发薪酬")
    add(None, None, "—研发人员股份支付", "data", "研发股份")
    add(None, None, "—研发人员数量", "data", "研发人数")
    add(None, None, "—研发人员工资", "wage", "研发")
    # 供应商
    add("可变成本", "营业成本-工人薪酬-工厂折旧摊销-制造其他[-售后]", "供应商（含委外加工）", "supplier", None, key="供应商")
    # 银行 (财务费用 + 金融子公司利息支出/手续费支出, 后两者无金融业务时为0)
    add("固定成本", "财务费用+金融子公司利息及手续费支出", "银行", "bank", None, key="银行")
    # 政府
    add("固定成本", "不含增值税", "政府", "sum_gov", None, key="政府")
    add(None, None, "—税金及附加", "ref", R["税金及附加"])
    add(None, None, "—所得税费用", "ref", R["所得税费用"])
    # 折旧
    add("固定成本", "固定资产+投资性房地产折旧", "折旧", "dep", None, key="折旧")
    # 摊销
    add("固定成本", None, "摊销", "sum_amort", None, key="摊销")
    add(None, None, "—无形资产摊销", "ref", R.get("无形资产摊销"))
    add(None, None, "—长期待摊费用摊销", "ref", R.get("长期待摊费用摊销"))
    add(None, None, "—使用权资产折旧与摊销", "ref", R.get("使用权资产折旧"))
    # 净利润
    add(None, "少数股东损益", "净利润-少数股东", "ref", R["少数股东损益"], key="净利润-少数股东")
    add(None, "归母净利润", "净利润-母公司股东", "ref", R["净利润母公司"], key="净利润-母公司股东")
    detail_last = r - 1

    # 附: 间接费用折旧摊销 + 供应商口径对照(营业成本口径供应商 vs residual 的差)
    r += 2
    r_fu = r
    add(None, None, "附：间接费用中的折旧摊销", "label", None)
    add(None, None, "—销售费用里的折旧摊销", "data", "销售折摊")
    rDA_sales = r - 1
    add(None, None, "—管理费用里的折旧摊销", "data", "管理折摊")
    rDA_mgmt = r - 1
    add(None, None, "—研发费用里的折旧摊销", "data", "研发折摊")
    rDA_rd = r - 1
    r_indirect = r
    add(None, None, "间接费用中的折旧摊销合计", "sum_da", (rDA_sales, rDA_rd))
    add(None, "营业成本-制造人工-工厂折摊-制造费用其他[-售后]", "营业成本口径供应商（参考）", "supplier_buildup", None, key="供应商参考")
    add(None, "=供应商(residual)-营业成本口径; 成本侧口径残差(已并入供应商)", "口径残差（已并入供应商）", "supplier_gap", None, key="口径残差")

    # ---- 写入细项拆分区 ----
    for (row, A, B, C, kind, payload) in layout:
        if A:
            ws.cell(row, 1, A)
        if B:
            ws.cell(row, 2, B)
        ws.cell(row, 3, C)
        for y in years:
            cc = yc[y]; L = CL(cc); d = nd.get(y, {})
            f = None
            if kind == "rev":
                f = "=" + ref(y, payload)
            elif kind == "grand":
                # 加总=各分配桥接科目之和(供应商为residual故恒=收入)
                parts = [dr[k] for k in ["经营损失","非经营损失","业务费用","售后服务费","员工薪酬",
                                          "供应商","银行","政府","折旧","摊销","净利润-少数股东","净利润-母公司股东"]]
                f = "=" + "+".join(f"{L}{p}" for p in parts)
            elif kind == "negsum":
                f = f"=-SUM({L}{op_first}:{L}{op_last})"
            elif kind == "nonop":
                f = f"=-({ref(y,R['营业外收入'])}-{ref(y,R['营业外支出'])})"
            elif kind == "biz":
                f = f"={L}{dr['销售相关费用']}+{L}{dr['管理相关费用']}+{L}{dr['研发相关费用']}+{L}{dr['制造其他']}"
            elif kind == "sum_sales":
                f = f"=SUM({L}{sales_first}:{L}{sales_last})" if sales else 0
            elif kind == "sum_mgmt":
                f = f"=SUM({L}{mgmt_first}:{L}{mgmt_last})" if mgmt else 0
            elif kind == "sum_rd":
                f = f"=SUM({L}{rd_first}:{L}{rd_last})" if rd else 0
            elif kind == "ref":
                f = ("=" + ref(y, payload)) if payload else 0
            elif kind == "bank":
                parts = [ref(y, R["财务费用"])]
                if R.get("利息支出"):
                    parts.append(ref(y, R["利息支出"]))
                if R.get("手续费及佣金支出"):
                    parts.append(ref(y, R["手续费及佣金支出"]))
                f = "=" + "+".join(parts)
            elif kind == "emp":
                pc = prevcol(y)
                shares = f"+{L}{dr['销售人员']+1}+{L}{dr['管理人员']+1}+{L}{dr['研发人员']+1}"
                # 销售/管理/研发股份支付 行 = r_*emp+2
                shares = f"+{L}{r_semp+2}+{L}{r_memp+2}+{L}{r_remp+2}"
                if pc:
                    f = (f"=IFERROR(({src}!{sc[y]}{R['应付职工薪酬']}-{src}!{pc}{R['应付职工薪酬']}"
                         f"+{src}!{sc[y]}{R['支付职工现金']})/{u},0){shares}")
                else:
                    f = f"=IFERROR({src}!{sc[y]}{R['支付职工现金']}/{u},0){shares}"
            elif kind == "mfgworker":
                f = f"={L}{dr['员工薪酬']}-{L}{r_semp}-{L}{r_memp}-{L}{r_remp}"
            elif kind == "sum2":
                f = f"=SUM({L}{row+1}:{L}{row+2})"
            elif kind == "wage":
                comp_map = {'制造': '制造工人', '销售': '销售人员', '管理': '管理人员', '研发': '研发人员'}
                comp_r = dr[comp_map[payload]]; hc_r = row - 1
                f = f'=IFERROR({L}{comp_r}*10000/{L}{hc_r},"")'  # 万元/人(比亚迪口径)
            elif kind == "supplier":
                # 供应商=营业总收入-其余各已披露分配项(residual): 营业成本这个上一级的"其他", 成本侧口径残差自然落此, 使加总≡收入
                parts = [dr[k] for k in ["经营损失","非经营损失","业务费用","售后服务费","员工薪酬",
                                          "银行","政府","折旧","摊销","净利润-少数股东","净利润-母公司股东"]]
                f = f"={L}{dr['收入']}-(" + "+".join(f"{L}{p}" for p in parts) + ")"
            elif kind == "supplier_buildup":
                # 参考: 营业成本口径供应商=营业成本-制造工人-(折旧+摊销-间接折摊)-制造其他[-售后(COGS年)]
                base_f = (f"={ref(y,R['营业成本'])}-{L}{dr['制造工人']}"
                          f"-({L}{dr['折旧']}+{L}{dr['摊销']}-{L}{r_indirect})-{L}{dr['制造其他']}")
                if y in wcogs:
                    base_f += f"-{L}{dr['售后服务费']}"
                f = base_f
            elif kind == "supplier_gap":
                f = f"={L}{dr['供应商']}-{L}{dr['供应商参考']}"
            elif kind == "sum_gov":
                f = f"=SUM({L}{row+1}:{L}{row+2})"
            elif kind == "dep":
                parts = [ref(y, R['固定资产折旧'])]
                if R.get('投资性房地产折旧'):
                    parts.append(ref(y, R['投资性房地产折旧']))
                f = "=" + "+".join(parts)
            elif kind == "sum_amort":
                f = f"=SUM({L}{row+1}:{L}{row+3})"
            elif kind == "sum_da":
                a, b = payload
                f = f"=SUM({L}{a}:{L}{b})"
            elif kind == "label":
                f = None
            elif kind == "data":
                wsr = cfg.get("ws_rows", {})
                if payload in wsr:   # 链接到附注底稿
                    wscol = cfg["ws_col"][y]
                    f = f"=IFERROR('{cfg['ws_sheet']}'!{wscol}{wsr[payload]},0)"
                elif payload == "制造其他" and cfg.get("mfg_other_pct"):
                    f = f"=ROUND({CL(cc)}{dr['收入']}*{cfg['mfg_other_pct']},2)"
                else:
                    v = d.get(payload)
                    f = v if v is not None else None
            if f is not None:
                ws.cell(row, cc, f)
            cell = ws.cell(row, cc)
            if kind == "wage":
                cell.number_format = "#,##0.0"
            elif kind == "data" and isinstance(payload, str) and "人数" in payload:
                cell.number_format = "#,##0"
            else:
                cell.number_format = AMT
        # 样式
        c3 = ws.cell(row, 3)
        if kind in ("rev", "grand", "biz", "emp", "supplier", "政府", "sum_gov", "dep", "sum_amort"):
            c3.font = BOLD

    # 年份表头 (细项区)
    ws.cell(38, 3, "项目数额")
    for y in years:
        ws.cell(38, yc[y], f"{y}年").font = BOLD
    ws.cell(37, 3, "细项拆分").font = BOLD

    # ============ 段1 总项统计 (VLOOKUP by label) ============
    bridge = ["收入", "加总", "经营损失", "非经营损失", "业务费用", "售后服务费", "员工薪酬",
              "供应商", "银行", "政府", "折旧", "摊销", "净利润-少数股东", "净利润-母公司股东"]
    blabel = {"加总": "下面各项加总", "售后服务费": "售后服务费（给客户预备的钱，原隶属于销售费用）",
              "业务费用": "业务费用（给广泛第三方的钱）", "供应商": "供应商（含委外加工）"}
    ws.cell(1, 1, "跳转：总项统计"); ws.cell(1, 2, f"单位：亿元")
    ws.cell(2, 3, title).font = BOLD
    for y in years:
        ws.cell(2, yc[y], f"{y}年").font = BOLD
    ws.cell(3, 3, "总项统计").font = BOLD
    ws.cell(4, 3, "项目数额").font = BOLD
    for y in years:
        ws.cell(4, yc[y], f"{y}年").font = BOLD
    rng = f"$C${base}:${CL(cols[-1])}${detail_last}"
    for i, k in enumerate(bridge):
        rr = 5 + i
        lab = blabel.get(k, k)
        ws.cell(rr, 3, lab)
        if k in ("rev", "加总", "经营损失", "非经营损失") or k == "收入":
            ws.cell(rr, 3).font = BOLD if k in ("收入", "加总") else Font()
        for y in years:
            cc = yc[y]
            ws.cell(rr, cc, f'=VLOOKUP($C{rr},{rng},{cc-2},FALSE())')
            ws.cell(rr, cc).number_format = AMT

    # ============ 段2 项目占比 ============
    ws.cell(20, 3, "项目占比").font = BOLD
    for y in years:
        ws.cell(20, yc[y], f"{y}年").font = BOLD
    rev_row1 = 5
    for i, k in enumerate(bridge):
        rr = 21 + i
        src_r = 5 + i
        lab = blabel.get(k, k)
        ws.cell(rr, 3, lab)
        if k == "收入":
            for y in years:
                ws.cell(rr, yc[y], 1).number_format = PCT
        elif k == "加总":
            last_ratio = 21 + len(bridge) - 1
            for y in years:
                L = CL(yc[y])
                ws.cell(rr, yc[y], f"=SUM({L}23:{L}{last_ratio})").number_format = PCT
        else:
            for y in years:
                L = CL(yc[y])
                ws.cell(rr, yc[y], f"={L}{src_r}/{L}${rev_row1}").number_format = PCT

    # ============ 段4 占比明细 (镜像细项) ============
    g = r + 2
    g_hdr = g
    ws.cell(g, 3, "细项占比").font = BOLD
    for y in years:
        ws.cell(g, yc[y], f"{y}年").font = BOLD
    g += 1
    s4_first = g
    rev_detail = dr["收入"]
    for (row, A, B, C, kind, payload) in layout:
        # 跳过非货币行: 标题、人均工资、人数(占比无意义)
        if kind in ("label", "wage"):
            continue
        if kind == "data" and isinstance(payload, str) and "人数" in payload:
            continue
        ws.cell(g, 3, C)
        for y in years:
            L = CL(yc[y])
            ws.cell(g, yc[y], f"={L}{row}/{L}${rev_detail}").number_format = PCT
        g += 1
    s4_last = g - 1

    # ---- 统一格式：参考公司深度报告模板"表2"（深蓝表头2E3160白字、思源黑体、三线表） ----
    NAVY = "2E3160"; FN = "思源黑体 CN"; FNM = "思源黑体 CN Medium"
    navy_fill = PatternFill("solid", fgColor=NAVY)
    light_fill = PatternFill("solid", fgColor="E7E9F0")
    white_b = Font(name=FNM, bold=True, color="FFFFFF", size=9)
    navy_b = Font(name=FNM, bold=True, color=NAVY, size=9)
    body = Font(name=FN, size=9); totb = Font(name=FNM, bold=True, size=9)
    grey_font = Font(name=FN, color="808080", size=9)
    src_font = Font(name=FN, color="404040", size=8)
    CEN = Alignment(horizontal="center"); RIGHT = Alignment(horizontal="right"); LEFT = Alignment(horizontal="left")
    med = Side(style="medium", color=NAVY); thin = Side(style="thin", color="BFBFBF")
    lastcol = cols[-1]
    last_row = s4_last
    # 1) 基础：全表思源黑体, 标签左对齐/数值右对齐
    for rr in range(2, last_row + 1):
        ws.cell(rr, 3).font = body; ws.cell(rr, 3).alignment = LEFT
        for y in years:
            cell = ws.cell(rr, yc[y]); cell.font = body; cell.alignment = RIGHT
    # 2) 表头band(深蓝白字): 标题(2)、总项数额(4)、占比(20)、细项数额(38)、占比明细(g_hdr)
    def band(rr):
        for c in range(3, lastcol + 1):
            cell = ws.cell(rr, c); cell.fill = navy_fill; cell.font = white_b
            cell.alignment = CEN if c > 3 else LEFT
    for hr in [2, 4, 20, 38, g_hdr]:
        band(hr)
    # 3) 段标题(总项统计3/细项拆分37): 浅蓝填充+深蓝粗体
    for tr in [3, 37]:
        for c in range(3, lastcol + 1):
            ws.cell(tr, c).fill = light_fill
        ws.cell(tr, 3).font = navy_b
    # 4) 收入/加总行加粗(总项5/6, 细项39/40)
    for tr in [5, 6, 39, 40]:
        for c in range(3, lastcol + 1):
            ws.cell(tr, c).font = totb
    # 5) 占比区灰字(项目占比21-34, 占比明细)
    for rr in list(range(21, 35)) + list(range(s4_first, s4_last + 1)):
        for y in years:
            cell = ws.cell(rr, yc[y])
            if not cell.font.bold:
                cell.font = grey_font
    # 6) 三线表外框: 顶线(标题2上) + 表头下线(4下) + 底线(last_row下)
    for c in range(3, lastcol + 1):
        ws.cell(2, c).border = Border(top=med)
        ws.cell(4, c).border = Border(bottom=thin)
        ws.cell(last_row, c).border = Border(bottom=med)
    # 7) 数据来源(表下方, 报告风)
    src = ws.cell(last_row + 1, 3, '数据来源：公司历年年度报告及三大财务报表、广发证券发展研究中心')
    src.font = src_font
    ws.merge_cells(start_row=last_row + 1, start_column=3, end_row=last_row + 1, end_column=lastcol)
    # 8) 层次标注：二级桥接科目(浅蓝底纹+深蓝加粗) / 三级小计(加粗)
    l2fill = PatternFill("solid", fgColor="E7E9F0")
    l2font = Font(name=FNM, bold=True, color=NAVY, size=9)
    l3font = Font(name=FNM, bold=True, size=9)
    L2 = ["经营损失", "非经营损失", "业务费用", "售后服务费", "员工薪酬", "供应商", "银行",
          "政府", "折旧", "摊销", "净利润-少数股东", "净利润-母公司股东"]
    L3 = ["销售相关费用", "管理相关费用", "研发相关费用", "制造工人", "销售人员", "管理人员", "研发人员"]
    for k in L2:
        rr = dr[k]
        for c in range(3, lastcol + 1):
            cell = ws.cell(rr, c); cell.fill = l2fill; cell.font = l2font
            cell.alignment = LEFT if c == 3 else RIGHT
    for k in L3:
        rr = dr[k]
        for c in range(3, lastcol + 1):
            cell = ws.cell(rr, c); cell.font = l3font
            cell.alignment = LEFT if c == 3 else RIGHT
    # 总项统计二级(7-18)同步标注
    for rr in range(7, 19):
        for c in range(3, lastcol + 1):
            cell = ws.cell(rr, c); cell.fill = l2fill; cell.font = l2font
            cell.alignment = LEFT if c == 3 else RIGHT
    # A1/A2 内部导航
    ws['A1'].value = '跳转：总项统计'; ws['A1'].hyperlink = f"#'{name}'!C3"; ws['A1'].font = Font(name=FN, color="0563C1", underline="single", size=9)
    ws['A2'].value = '跳转：细项拆分'; ws['A2'].hyperlink = f"#'{name}'!C37"; ws['A2'].font = Font(name=FN, color="0563C1", underline="single", size=9)
    # 列宽
    ws.column_dimensions['A'].width = 11
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 36
    for y in years:
        ws.column_dimensions[CL(yc[y])].width = 11
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D5"
    # 隐藏诊断行(营业成本口径供应商参考 / 口径残差)——含占比明细镜像; 需要时可取消隐藏
    for rr in range(1, s4_last + 2):
        if ws.cell(rr, 3).value in ('营业成本口径供应商（参考）', '口径残差（已并入供应商）'):
            ws.row_dimensions[rr].hidden = True
    return ws, dr, detail_last
