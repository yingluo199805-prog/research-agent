# -*- coding: utf-8 -*-
"""把数据端导出的『利润表/资产负债表/现金流量表』三个单表合并为一个公司工作簿的『XX三表』表。
- 年份列改为降序(最新在左, B列起), 单位百万元, 标签在A列
- 段序: 利润表 / 资产负债表 / 现金流量表
返回每段起始行, 便于确定 src_rows。
"""
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as CL

def load_stmt(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    # 年份行=row3, 数据列 B.. ; 升序 -> 取年份与列
    yrs = []
    for c in range(2, ws.max_column + 1):
        v = ws.cell(3, c).value
        if v and '年报' in str(v):
            yrs.append((int(str(v)[:4]), c))
    rows = []  # (label, {year:value})
    for r in range(4, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if lab is None or str(lab).strip() == '':
            continue
        d = {y: ws.cell(r, c).value for y, c in yrs}
        rows.append((str(lab), d))
    years = sorted([y for y, _ in yrs], reverse=True)
    return years, rows

def build_company_sheet(out_path, sheet_name, stmt_paths, unit_label='（单位：百万元，CNY）'):
    """stmt_paths: dict 利润表/资产负债表/现金流量表 -> file path"""
    all_years = None
    sections = []
    for title, key in [('一、利润表', '利润表'), ('二、资产负债表', '资产负债表'), ('三、现金流量表', '现金流量表')]:
        years, rows = load_stmt(stmt_paths[key])
        all_years = years if all_years is None else all_years
        sections.append((title, rows))
    import os
    if os.path.exists(out_path):
        wb = openpyxl.load_workbook(out_path)
    else:
        wb = openpyxl.Workbook(); wb.remove(wb.active)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 0)
    ycol = {y: 2 + i for i, y in enumerate(all_years)}  # B=最新
    section_start = {}
    r = 1
    for title, rows in sections:
        ws.cell(r, 1, title).font = Font(bold=True); r += 1
        ws.cell(r, 1, unit_label); r += 1
        for y in all_years:
            ws.cell(r, ycol[y], f'{y} 年报').font = Font(bold=True)
        r += 1
        section_start[title] = r
        for lab, d in rows:
            ws.cell(r, 1, lab)
            for y in all_years:
                v = d.get(y)
                ws.cell(r, ycol[y], v if v not in ('--', None) else None)
            r += 1
        r += 2
    ws.column_dimensions['A'].width = 34
    for y in all_years:
        ws.column_dimensions[CL(ycol[y])].width = 12
    wb.save(out_path)
    return all_years, section_start

if __name__ == '__main__':
    """用法: python3 combine_three_statements.py <三表文件夹> <输出sheet名> <输出xlsx>
    例:   python3 combine_three_statements.py 华域汽车2015-2025三表 华域汽车三表 华域汽车A股.xlsx
    相对路径均相对 REVENUE_SPLIT_BASE（默认 <仓库>/outputs/收入分配）。文件夹内须有
    *-利润表-*.xlsx / *-资产负债表-*.xlsx / *-现金流量表-*.xlsx 三个数据端导出文件。"""
    import sys, os, glob
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import common
    if len(sys.argv) < 4:
        raise SystemExit(__doc__ if __doc__ else "need 3 args")
    folder, sheet, outfile = sys.argv[1:4]
    folder = folder if os.path.isabs(folder) else common.base_path(folder)
    outfile = outfile if os.path.isabs(outfile) else common.base_path(outfile)
    paths = {}
    for s in ['利润表', '资产负债表', '现金流量表']:
        hit = glob.glob(os.path.join(folder, f"*-{s}-*.xlsx"))
        if not hit:
            raise SystemExit(f"{folder} 缺 {s} 导出文件")
        paths[s] = hit[0]
    years, ss = build_company_sheet(outfile, sheet, paths)
    print('out:', outfile)
    print('years:', years)
    print('section_start:', ss)
