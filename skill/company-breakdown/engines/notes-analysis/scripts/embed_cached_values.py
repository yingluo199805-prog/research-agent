# -*- coding: utf-8 -*-
"""把 formulas 重算副本（*_recalc.xlsx，只有值）里的结果回写成原工作簿公式单元格的缓存值（<v>），
原文件的公式、样式、合并单元格、图片全部保留。这样手机/网页预览、Numbers、未开自动重算的 Excel 也能直接看到数。

用法: python3 embed_cached_values.py <原.xlsx> [<重算副本.xlsx>]   （副本默认 = 原名 + _recalc.xlsx）
做法: 解压原 xlsx → 按 workbook.xml/rels 把 sheet 名映射到 sheetN.xml → 正则找 <c r=.. ><f>..</f><v></v></c> →
      从副本(openpyxl data_only)取该单元格值写进 <v>（字符串加 t="str"、布尔 t="b"、错误 t="e"）→ 用 zipfile 原顺序重打包。
      同时保留 workbook.xml 的 fullCalcOnLoad="1"，Excel 打开仍会全量重算，缓存值只是预览用。
"""
import sys, os, re, zipfile, shutil, tempfile
import openpyxl

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def sheet_map(z):
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid2target = dict(re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    rid2target.update({m[1]: m[0] for m in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        t = rid2target[rid]
        out[name.replace("&amp;", "&")] = "xl/" + t if not t.startswith("/") else t.lstrip("/")
    return out


def fmt_value(v):
    if v is None:
        return None, None
    if isinstance(v, bool):
        return ("1" if v else "0"), "b"
    if isinstance(v, (int, float)):
        if v != v or v in (float("inf"), float("-inf")):
            return "#NUM!", "e"
        return repr(float(v)) if isinstance(v, float) else str(v), None
    s = str(v)
    if s.startswith("#"):
        return s, "e"
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return s, "str"


def embed(src, recalc=None):
    recalc = recalc or os.path.splitext(src)[0] + "_recalc.xlsx"
    vals = openpyxl.load_workbook(recalc, data_only=True)
    zin = zipfile.ZipFile(src)
    smap = sheet_map(zin)
    tmp = src + ".tmp"
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    n_done = n_missing = 0
    cell_pat = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v></v></c>', re.S)
    for item in zin.infolist():
        data = zin.read(item.filename)
        sheetname = next((n for n, p in smap.items() if p == item.filename), None)
        if sheetname and sheetname in vals.sheetnames:
            ws = vals[sheetname]
            xml = data.decode("utf-8")

            def sub(m):
                nonlocal n_done, n_missing
                ref, attrs, f = m.group(1), m.group(2), m.group(3)
                v, t = fmt_value(ws[ref].value)
                if v is None:
                    n_missing += 1
                    return m.group(0)
                n_done += 1
                attrs2 = re.sub(r'\s+t="[^"]*"', "", attrs)
                if t:
                    attrs2 += f' t="{t}"'
                return f'<c r="{ref}"{attrs2}><f>{f}</f><v>{v}</v></c>'
            xml = cell_pat.sub(sub, xml)
            data = xml.encode("utf-8")
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src)
    print(f"embedded cached values into {os.path.basename(src)}: {n_done} formula cells filled, {n_missing} left empty")
    return src


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    embed(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
